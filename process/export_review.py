"""목업 1차(키워드) + 2차(채널 순회) 결과를 사람이 검수하기 쉬운 표로 내보낸다.

출력 (review/):
  목업_검수목록.csv    엑셀에서 바로 열리는 UTF-8 BOM CSV
  목업_검수목록.xlsx   시트: 검수목록(링크·필터·드롭다운), 채널수율, 긴영상_타임스탬프, 컬럼설명

사용: .venv/bin/python process/export_review.py
먼저 mock_analyze.py, mock_analyze_v2.py를 돌려 *_labeled.csv를 최신으로 만들어 둘 것.
"""
import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "review"

REVIEW_COLS = {  # 검수 칸: 이름 → 드롭다운 값 (None이면 자유 입력)
    "검수_관련여부": "O,X,?",
    "검수_원곡": None,
    "검수_차용유형": "R0 원형,R1 편곡,L1 가사전체,L2 가사부분,L3 개사,M1 선율전체,M2 선율부분,M3 샘플링,S1 토리,C1 컨셉,C2 제목만",
    "검수_갈래": "T 전통,F 퓨전,M 현대,U 아마추어,AI 생성",
    "검수_보컬": "기악,가사보컬,구음스캣,랩포함",
    "검수_편성": "전통반주,국악앙상블,관현악,밴드,국악+밴드,전자,서양실내악,합창",
    "검수_오디오등급": "1,2,3,4,5",
    "검수_메모": None,
}

GUIDE = [
    ("차수", "v1 = 1차 키워드 검색(곡명+명창/퓨전국악), v2 = 2차 채널·시리즈 순회"),
    ("수집경로", "v1: 검색어 / v2: 시드 채널 이름"),
    ("매칭곡", "자동으로 찾은 anchor 민요 (total_metadata 곡명 기준). 여러 곡이면 | 로 구분"),
    ("매칭근거", "paren = 제목 괄호 속 원곡명 (가장 확실), title = 제목에 곡명, (v1) 가사구절 = 자막에서 곡 고유 구절"),
    ("제목패턴", "fm_live, paren(신곡명(원곡) 또는 셋리스트), artist-song(아티스트 - 곡), keyword(곡명만 포함)"),
    ("갈래추정", "규칙 기반: T 전통 / F 퓨전 / M 현대·미디어 / U 커버·강습 / ? 미상"),
    ("영상형태", "topic_audio 공식음원, official_mv, broadcast 방송, live_stage 공연, lesson 강습, short 숏츠"),
    ("오디오등급", "1 공식 음원 / 2 방송 스튜디오 라이브 / 3 MV·라이브클립·방송무대 / 4 공연 실황 / 5 직캠·재업로드 (자동 추정)"),
    ("AI의심", "설명란·제목에 AI 생성 고지가 있는 경우"),
    ("여러곡", "제목에 anchor 곡이 2개 이상 → 메들리·셋리스트, 구간 분할 필요"),
    ("가사구절(v1)", "v1에서 자막을 받은 영상만. 곡 고유 구절이 잡힌 곡(잡힌 수)"),
    ("검수_*", "사람이 채우는 칸. 엑셀에서 드롭다운으로 선택. 차용유형 코드는 BRAINSTORM 문서 §3 참고"),
]


def fmt_dur(sec):
    if pd.isna(sec):
        return ""
    sec = int(sec)
    return f"{sec // 3600}:{sec % 3600 // 60:02d}:{sec % 60:02d}" if sec >= 3600 else f"{sec // 60}:{sec % 60:02d}"


def load_v1():
    d = pd.read_csv(ROOT / "mock/mock_videos_labeled.csv")
    return pd.DataFrame({
        "차수": "v1", "수집경로": d.query_song + " (" + d.query_axis + ")", "채널": d.channel, "제목": d.title,
        "URL": d.url, "업로드연도": d.upload_year, "길이": d.duration.map(fmt_dur), "길이(초)": d.duration,
        "조회수": d.views, "매칭곡": d.title_songs.fillna(""),
        "매칭근거": d.lyric_top.notna().map({True: "title+가사구절", False: "title"}),
        "제목패턴": "", "갈래추정": d.spectrum_guess, "영상형태": d.media_form, "오디오등급": "",
        "AI의심": "", "여러곡": d.title_songs.fillna("").str.contains(r"\|"),
        "자동자막": d.sub_kind, "가사구절(v1)": d.lyric_songs.fillna(""), "video_id": d.video_id,
    })


def load_v2():
    d = pd.read_csv(ROOT / "mock_v2/mock_v2_labeled.csv")
    return pd.DataFrame({
        "차수": "v2", "수집경로": d.seed, "채널": d.seed, "제목": d.title, "URL": d.url,
        "업로드연도": d.upload_year, "길이": d.duration.map(fmt_dur), "길이(초)": d.duration, "조회수": d.views,
        "매칭곡": d.songs.fillna(""), "매칭근거": d.match_source, "제목패턴": d.pattern,
        "갈래추정": d.spectrum_guess, "영상형태": d.media_form, "오디오등급": d.audio_tier,
        "AI의심": d.ai_suspect.map({True: "의심", False: ""}), "여러곡": d.n_songs >= 2,
        "자동자막": d.has_ko_auto.map({1: "auto", 0: "none"}), "가사구절(v1)": "", "video_id": d.video_id,
    })


def channel_yield():
    con = sqlite3.connect(ROOT / "mock_v2/mock_v2.sqlite")
    cl = pd.read_sql("SELECT * FROM channel_list", con)
    seeds = pd.read_sql("SELECT * FROM seed_log", con).set_index("seed")
    cl["matched"] = cl.songs.fillna("") != ""
    g = cl.groupby(["seed_type", "seed"]).agg(전체영상=("video_id", "size"), 민요매칭=("matched", "sum"),
                                              괄호확정=("match_source", lambda x: (x == "paren").sum()))
    g["매칭비율"] = (g.민요매칭 / g.전체영상).round(3)
    g = g.join(seeds["seconds"].rename("목록수집_초"), on="seed").reset_index()
    return g.rename(columns={"seed_type": "시드유형", "seed": "시드"})


def style(path, n_rows):
    wb = load_workbook(path)
    ws = wb["검수목록"]
    header_fill = PatternFill("solid", fgColor="DDE4EE")
    review_fill = PatternFill("solid", fgColor="FFF4D6")
    headers = [c.value for c in ws[1]]
    widths = {"제목": 60, "URL": 30, "수집경로": 22, "채널": 22, "매칭곡": 24, "가사구절(v1)": 22, "검수_메모": 30,
              "검수_원곡": 16, "검수_차용유형": 16, "검수_편성": 14}
    for i, h in enumerate(headers, 1):
        col = get_column_letter(i)
        ws[f"{col}1"].font = Font(bold=True)
        ws[f"{col}1"].fill = review_fill if h.startswith("검수_") else header_fill
        ws.column_dimensions[col].width = widths.get(h, 11)
        if h.startswith("검수_"):
            for r in range(2, n_rows + 2):
                ws[f"{col}{r}"].fill = review_fill
            if REVIEW_COLS.get(h):
                dv = DataValidation(type="list", formula1=f'"{REVIEW_COLS[h]}"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{col}2:{col}{n_rows + 1}")
        if h == "URL":
            for r in range(2, n_rows + 2):
                c = ws[f"{col}{r}"]
                c.hyperlink = c.value
                c.font = Font(color="0563C1", underline="single")
        if h == "제목":
            for r in range(2, n_rows + 2):
                ws[f"{col}{r}"].alignment = Alignment(wrap_text=False)
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions
    for name in wb.sheetnames[1:]:
        s = wb[name]
        for c in s[1]:
            c.font = Font(bold=True)
            c.fill = header_fill
        for i in range(1, s.max_column + 1):
            s.column_dimensions[get_column_letter(i)].width = 18
        s.freeze_panes = "A2"
    wb["컬럼설명"].column_dimensions["B"].width = 100
    wb.save(path)


def main():
    OUT.mkdir(exist_ok=True)
    df = pd.concat([load_v1(), load_v2()], ignore_index=True)
    df = df.drop_duplicates("video_id", keep="first")
    for c in REVIEW_COLS:
        df[c] = ""
    df.insert(0, "번호", range(1, len(df) + 1))
    df.to_csv(OUT / "목업_검수목록.csv", index=False, encoding="utf-8-sig")

    seg_path = ROOT / "mock/long/ftcCjBbMJHY.segments.csv"
    segs = pd.read_csv(seg_path) if seg_path.exists() else pd.DataFrame()
    if len(segs):
        segs.insert(0, "URL", "https://youtu.be/ftcCjBbMJHY")
        segs["시작"] = segs.start.map(fmt_dur)
        segs["끝"] = segs.end.map(fmt_dur)
        segs["바로가기"] = [f"https://youtu.be/ftcCjBbMJHY?t={int(s)}" for s in segs.start]
    xlsx = OUT / "목업_검수목록.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as w:
        df.to_excel(w, sheet_name="검수목록", index=False)
        channel_yield().to_excel(w, sheet_name="채널수율", index=False)
        segs.to_excel(w, sheet_name="긴영상_타임스탬프", index=False)
        pd.DataFrame(GUIDE, columns=["컬럼", "설명"]).to_excel(w, sheet_name="컬럼설명", index=False)
    style(xlsx, len(df))
    print(f"rows={len(df)} (v1 {sum(df.차수 == 'v1')}, v2 {sum(df.차수 == 'v2')}) -> {OUT}")


if __name__ == "__main__":
    main()
